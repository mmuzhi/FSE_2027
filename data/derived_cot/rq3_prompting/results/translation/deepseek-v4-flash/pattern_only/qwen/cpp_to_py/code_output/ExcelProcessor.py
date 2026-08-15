from openpyxl import Workbook, load_workbook


class ExcelProcessor:
    def read_excel(self, file_name):
        data = []
        try:
            wb = load_workbook(file_name)
            try:
                if "Sheet1" not in wb.sheetnames:
                    return []
                ws = wb["Sheet1"]

                for row_index in range(1, ws.max_row + 1):
                    row_data = []
                    for col_index in range(1, ws.max_column + 1):
                        value = ws.cell(row=row_index, column=col_index).value
                        if isinstance(value, int) and not isinstance(value, bool):
                            row_data.append(value)
                        elif isinstance(value, str):
                            row_data.append(value)
                    data.append(row_data)
            finally:
                wb.close()
        except Exception:
            return []
        return data

    def write_excel(self, data, file_name):
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"

            for row_index, row in enumerate(data, start=1):
                for col_index, value in enumerate(row, start=1):
                    ws.cell(row=row_index, column=col_index, value=value)

            wb.save(file_name)
            wb.close()
            return 1
        except Exception:
            return 0

    def process_excel_data(self, N, save_file_name):
        data = self.read_excel(save_file_name)
        if not data or N >= len(data[0]):
            return (0, "")

        new_data = []
        for row in data:
            new_row = list(row)
            value = row[N]

            if isinstance(value, str):
                if value == "":
                    new_row.append("")
                elif not ('0' <= value[0] <= '9'):
                    now = ''.join(
                        chr(ord(ch) - 32) if ord(ch) > 90 else ch
                        for ch in value
                    )
                    new_row.append(now)
                else:
                    new_row.append(value[0])
            else:
                new_row.append(value)

            new_data.append(new_row)

        dot = save_file_name.rfind('.')
        if dot == -1:
            new_file_name = save_file_name + "_process.xlsx"
        else:
            new_file_name = save_file_name[:dot] + "_process.xlsx"

        success = self.write_excel(new_data, new_file_name)
        return (success, new_file_name)