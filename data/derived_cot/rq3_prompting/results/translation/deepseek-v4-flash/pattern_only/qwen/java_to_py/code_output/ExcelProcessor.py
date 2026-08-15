from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import WINDOWS_EPOCH, to_excel
from datetime import datetime, time, timedelta


def _to_java_int(value):
    if isinstance(value, float):
        if value != value:  # NaN
            return 0
        if value >= 2147483647:
            return 2147483647
        if value <= -2147483648:
            return -2147483648
    return int(value)


def readExcel(fileName):
    try:
        wb = load_workbook(fileName)
        ws = wb.worksheets[0]
    except OSError:
        return None

    data = []
    if not ws._cells:
        return data

    offset = (getattr(wb, 'epoch', WINDOWS_EPOCH) - WINDOWS_EPOCH).days

    row_indices = sorted({r for r, c in ws._cells.keys()})
    for row_idx in row_indices:
        row_cells = [cell for (r, c), cell in ws._cells.items() if r == row_idx]
        row_cells.sort(key=lambda cell: cell.column)
        rowData = []
        for cell in row_cells:
            cell_type = cell.data_type
            value = cell.value
            if cell_type == 's':
                rowData.append(value)
            elif cell_type in ('n', 'd') and value is not None:
                if cell_type == 'd':
                    if isinstance(value, datetime):
                        val = to_excel(value, offset=offset)
                    elif isinstance(value, time):
                        val = (value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6) / 86400.0
                    elif isinstance(value, timedelta):
                        val = value.total_seconds() / 86400.0
                    else:
                        val = value
                else:
                    val = value
                rowData.append(_to_java_int(val))
            else:
                rowData.append(None)
        data.append(rowData)
    return data


def writeExcel(data, fileName):
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        for i, rowData in enumerate(data):
            for j, value in enumerate(rowData):
                if isinstance(value, str):
                    ws.cell(row=i + 1, column=j + 1, value=value)
                elif type(value) is int and -2147483648 <= value <= 2147483647:
                    ws.cell(row=i + 1, column=j + 1, value=value)
        wb.save(fileName)
        return True
    except OSError:
        return False


def processExcelData(N, saveFileName):
    data = readExcel(saveFileName)
    if data is None or N >= len(data[0]):
        return None
    if N < 0:
        raise IndexError("Index out of bounds")
    newData = []
    for row in data:
        newRow = list(row)
        value = row[N]
        if isinstance(value, str):
            newRow.append(value.upper())
        else:
            newRow.append(value)
        newData.append(newRow)
    newFileName = saveFileName.replace(".xlsx", "_process.xlsx")
    success = writeExcel(newData, newFileName)
    return newFileName if success else None