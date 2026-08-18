from typing import List, Tuple, Union
from openpyxl import Workbook, load_workbook

Cell = Union[str, int]


class ExcelProcessor:

    def read_excel(self, file_name: str) -> List[List[Cell]]:
        data: List[List[Cell]] = []
        try:
            wb = load_workbook(file_name)
            sheet = wb["Sheet1"]

            row_count = sheet.max_row
            col_count = sheet.max_column

            for row_index in range(1, row_count + 1):
                row_data: List[Cell] = []
                for col_index in range(1, col_count + 1):
                    value = sheet.cell(row=row_index, column=col_index).value
                    # Only Integer and String cell types are kept (bool excluded),
                    # mirroring the XLValueType checks.
                    if isinstance(value, bool):
                        continue
                    elif isinstance(value, int):
                        row_data.append(value)
                    elif isinstance(value, str):
                        row_data.append(value)
                data.append(row_data)
            wb.close()
        except Exception:
            return []
        return data

    def write_excel(self, data: List[List[Cell]], file_name: str) -> int:
        try:
            wb = Workbook()
            sheet = wb.active
            sheet.title = "Sheet1"

            for row_index in range(len(data)):
                for col_index in range(len(data[row_index])):
                    value = data[row_index][col_index]
                    if isinstance(value, int) and not isinstance(value, bool):
                        sheet.cell(row=row_index + 1, column=col_index + 1, value=value)
                    else:
                        sheet.cell(row=row_index + 1, column=col_index + 1, value=value)
            wb.save(file_name)
            wb.close()
            return 1
        except Exception:
            return 0

    def process_excel_data(self, N: int, save_file_name: str) -> Tuple[int, str]:
        data = self.read_excel(save_file_name)
        if not data or N >= len(data[0]):
            return (0, "")

        new_data: List[List[Cell]] = []
        for row in data:
            new_row: List[Cell] = list(row)
            value = row[N]
            if isinstance(value, str):
                if not ('0' <= value[0] <= '9'):
                    # Exact char arithmetic as in C++: any char > 'Z' is shifted by -32
                    now = list(value)
                    for i in range(len(now)):
                        if ord(now[i]) > ord('Z'):
                            now[i] = chr(ord(now[i]) - 32)
                    new_row.append(''.join(now))
                else:
                    new_row.append(value[0:1])
            else:
                new_row.append(value)
            new_data.append(new_row)

        # find_last_of('.'): if no '.', npos -> whole string is kept
        idx = save_file_name.rfind('.')
        base = save_file_name if idx == -1 else save_file_name[:idx]
        new_file_name = base + "_process.xlsx"
        success = self.write_excel(new_data, new_file_name)
        return (success, new_file_name)