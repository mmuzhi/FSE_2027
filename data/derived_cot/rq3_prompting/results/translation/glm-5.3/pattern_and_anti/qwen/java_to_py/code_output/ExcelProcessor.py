from typing import Any, List, Optional

from openpyxl import Workbook, load_workbook


class ExcelProcessor:

    def __init__(self):
        pass

    def read_excel(self, file_name: str) -> Optional[List[List[Any]]]:
        try:
            workbook = load_workbook(file_name)
            sheet = workbook.worksheets[0]
            data = []
            for row in sheet.iter_rows():
                row_data = []
                for cell in row:
                    value = cell.value
                    if isinstance(value, str):
                        row_data.append(value)
                    elif isinstance(value, (int, float)) and not isinstance(value, bool):
                        # Java: (int) cast truncates toward zero; int() does the same
                        row_data.append(int(value))
                    else:
                        row_data.append(None)
                data.append(row_data)
            workbook.close()
            return data
        except OSError:
            # Java catches IOException (e.g. missing file) and returns null
            return None

    def write_excel(self, data: List[List[Any]], file_name: str) -> bool:
        try:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            for i, row_data in enumerate(data):
                for j, value in enumerate(row_data):
                    cell = sheet.cell(row=i + 1, column=j + 1)
                    # Java only sets String and Integer; other cells stay empty
                    if isinstance(value, str):
                        cell.value = value
                    elif isinstance(value, int) and not isinstance(value, bool):
                        cell.value = value
            workbook.save(file_name)
            return True
        except OSError:
            return False

    def process_excel_data(self, n: int, save_file_name: str) -> Optional[str]:
        data = self.read_excel(save_file_name)
        # Note: like Java, empty data raises an error on data[0] (IndexOutOfBounds/IndexError)
        if data is None or n >= len(data[0]):
            return None
        new_data = []
        for row in data:
            new_row = list(row)
            value = row[n]
            if isinstance(value, str):
                new_row.append(value.upper())
            else:
                new_row.append(value)
            new_data.append(new_row)
        new_file_name = save_file_name.replace(".xlsx", "_process.xlsx")
        success = self.write_excel(new_data, new_file_name)
        return new_file_name if success else None