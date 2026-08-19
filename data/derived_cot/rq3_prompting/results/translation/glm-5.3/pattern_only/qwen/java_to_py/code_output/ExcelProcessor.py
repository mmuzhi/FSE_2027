from openpyxl import Workbook, load_workbook


class ExcelProcessor:

    def read_excel(self, file_name):
        data = []
        workbook = None
        try:
            workbook = load_workbook(file_name, read_only=True)
            sheet = workbook.worksheets[0]
            for row in sheet.iter_rows(values_only=True):
                row_data = []
                for value in row:
                    if isinstance(value, str):
                        row_data.append(value)
                    elif isinstance(value, (int, float)) and not isinstance(value, bool):
                        # Java: (int) getNumericCellValue() -> truncating cast
                        row_data.append(int(value))
                    else:
                        row_data.append(None)
                data.append(row_data)
        except OSError:
            # IOException -> null
            return None
        finally:
            if workbook is not None:
                workbook.close()
        return data

    def write_excel(self, data, file_name):
        workbook = Workbook()
        try:
            sheet = workbook.active
            sheet.title = "Sheet1"
            for i, row_data in enumerate(data):
                for j, value in enumerate(row_data):
                    if isinstance(value, str):
                        sheet.cell(row=i + 1, column=j + 1, value=value)
                    elif isinstance(value, int) and not isinstance(value, bool):
                        sheet.cell(row=i + 1, column=j + 1, value=value)
                    # other types: cell left blank, same as Java (no setCellValue call)
            workbook.save(file_name)
            return True
        except OSError:
            return False
        finally:
            workbook.close()

    def process_excel_data(self, n, save_file_name):
        data = self.read_excel(save_file_name)
        if data is None or n >= len(data[0]):
            return None
        new_data = []
        for row in data:
            new_row = list(row)
            value = row[n]
            if isinstance(value, str):
                new_row.append(value.upper())
            else:
                new_row.append(value)
            new_data.append(new_row)
        new_file_name = save_file_name.replace(".xlsx", "_process.xlsx")
        success = self.write_excel(new_data, new_file_name)
        return new_file_name if success else None