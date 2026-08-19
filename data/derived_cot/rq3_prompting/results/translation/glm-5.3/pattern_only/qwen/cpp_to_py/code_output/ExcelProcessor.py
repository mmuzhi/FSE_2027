from typing import List, Tuple, Union

import openpyxl

Cell = Union[str, int]


class ExcelProcessor:

    def read_excel(self, file_name: str) -> List[List[Cell]]:
        data: List[List[Cell]] = []
        try:
            wb = openpyxl.load_workbook(file_name)
            sheet = wb["Sheet1"]

            row_count = sheet.max_row
            col_count = sheet.max_column

            for row_index in range(1, row_count + 1):
                row_data: List[Cell] = []
                for col_index in range(1, col_count + 1):
                    value = sheet.cell(row=row_index, column=col_index).value
                    if isinstance(value, int) and not isinstance(value, bool):
                        row_data.append(int(value))
                    elif isinstance(value, str):
                        row_data.append(value)
                data.append(row_data)
            wb.close()
        except Exception:
            return []
        return data

    def write_excel(self, data: List[List[Cell]], file_name: str) -> int:
        try:
            wb = openpyxl.Workbook()
            sheet = wb.active
            sheet.title = "Sheet1"

            for row_index, row in enumerate(data):
                for col_index, cell_value in enumerate(row):
                    sheet.cell(row=row_index + 1, column=col_index + 1,
                               value=cell_value)
            wb.save(file_name)
            wb.close()
            return 1
        except Exception:
            return 0

    def process_excel_data(self, N: int, save_file_name: str) -> Tuple[int, str]:
        data = self.read_excel(save_file_name)
        if not data or N >= len(data[0]):
            return (0, "")

        new_data: List[List[Cell]] = []
        for row in data:
            new_row = list(row)
            value = row[N]
            if isinstance(value, str):
                first = value[:1]
                if not ('0' <= first <= '9'):
                    now = ''.join(
                        chr(ord(c) - 32) if c > 'Z' else c
                        for c in value
                    )
                    new_row.append(now)
                else:
                    new_row.append(first)
            else:
                new_row.append(value)
            new_data.append(new_row)

        dot = save_file_name.rfind('.')
        base = save_file_name if dot == -1 else save_file_name[:dot]
        new_file_name = base + "_process.xlsx"
        success = self.write_excel(new_data, new_file_name)
        return (success, new_file_name)